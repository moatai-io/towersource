#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "python-docx>=1.1.2",
#   "pillow>=10.0",
# ]
# ///
"""md_docx_publisher.py — Markdown to polished, Google-Docs-safe DOCX."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PALETTES = {
    "publish": {
        "font": "Noto Sans CJK KR",
        "mono": "DejaVu Sans Mono",
        "primary": "111111",
        "accent": "111111",
        "text": "242424",
        "table_header_bg": "202020",
        "table_alt_bg": "F7F7F5",
        "table_border": "D8D5CF",
    },
    "hantoo": {
        "font": "Noto Sans CJK KR",
        "mono": "DejaVu Sans Mono",
        "primary": "161616",
        "accent": "161616",
        "text": "242424",
        "table_header_bg": "202020",
        "table_alt_bg": "F7F7F5",
        "table_border": "D8D5CF",
    },
}


def die(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def run(cmd: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(cmd, text=True, capture_output=True, check=False)


def require(cmd: str) -> str:
    found = shutil.which(cmd)
    if not found:
        die(f"'{cmd}' not found")
    return found


def rgb(hexstr: str) -> RGBColor:
    return RGBColor(int(hexstr[:2], 16), int(hexstr[2:4], 16), int(hexstr[4:], 16))


def render_mermaid(md: str, tmp: Path) -> str:
    mmdc = shutil.which("mmdc")
    if not mmdc:
        return md

    def repl(match: re.Match[str]) -> str:
        idx = len(list(tmp.glob("mermaid-*.png"))) + 1
        mmd = tmp / f"mermaid-{idx:02d}.mmd"
        png = tmp / f"mermaid-{idx:02d}.png"
        mmd.write_text(match.group(1).strip() + "\n", encoding="utf-8")
        pup = tmp / "puppeteer.json"
        chrome = shutil.which("chromium") or shutil.which("google-chrome")
        if chrome:
            pup.write_text('{"executablePath":"' + chrome + '","args":["--no-sandbox","--disable-dev-shm-usage"]}', encoding="utf-8")
        cmd = [mmdc, "-i", str(mmd), "-o", str(png), "-b", "white", "-w", "1800", "--scale", "2"]
        if pup.exists():
            cmd += ["-p", str(pup)]
        p = run(cmd)
        return f"\n![]({png})\n" if p.returncode == 0 and png.exists() else match.group(0)

    return re.sub(r"```mermaid\s*\n(.*?)```", repl, md, flags=re.S)


def find_font() -> str | None:
    for font in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    ):
        if Path(font).exists():
            return font
    return None


def looks_like_ascii_diagram(text: str) -> bool:
    lines = [x.rstrip() for x in text.splitlines() if x.strip()]
    if len(lines) < 2:
        return False
    diagram_chars = sum(sum(ch in "+-|/\\<>[]{}┌┐└┘─│├┤┬┴┼→←↑↓" for ch in line) for line in lines)
    return diagram_chars >= 8 and diagram_chars / max(sum(map(len, lines)), 1) > 0.12


def render_ascii_diagrams(md: str, tmp: Path) -> str:
    font_path = find_font()
    font = ImageFont.truetype(font_path, 22) if font_path else ImageFont.load_default()

    def repl(match: re.Match[str]) -> str:
        lang, body = match.group(1).lower(), match.group(2).strip("\n")
        if lang not in {"ascii", "diagram", "textdiagram", "txt"} and not (not lang and looks_like_ascii_diagram(body)):
            return match.group(0)
        idx = len(list(tmp.glob("ascii-*.png"))) + 1
        png = tmp / f"ascii-{idx:02d}.png"
        lines = body.splitlines() or [""]
        left, top, pad = 26, 24, 26
        widths = [ImageDraw.Draw(Image.new("RGB", (1, 1))).textbbox((0, 0), line, font=font)[2] for line in lines]
        line_h = int(font.size * 1.55)
        img = Image.new("RGB", (max(widths, default=1) + left + pad, line_h * len(lines) + top + pad), "white")
        draw = ImageDraw.Draw(img)
        draw.rounded_rectangle([4, 4, img.width - 5, img.height - 5], radius=12, outline="#D6D3CC", width=2, fill="#FBFAF7")
        for i, line in enumerate(lines):
            draw.text((left, top + i * line_h), line, font=font, fill="#1E1E1E")
        img.save(png)
        return f"\n![]({png})\n"

    return re.sub(r"```([A-Za-z0-9_-]*)\s*\n(.*?)```", repl, md, flags=re.S)


def normalize_md(md: str) -> str:
    # ponytail: tiny normalization only; use a Markdown AST pass if semantic rewrites become necessary.
    md = re.sub(r"\n{3,}", "\n\n", md)
    return re.sub(r"<!--\s*pagebreak\s*-->", "\\newpage", md, flags=re.I)


def optimize_images(md: str, base: Path, tmp: Path, max_width: int, quality: int) -> str:
    if max_width <= 0:
        return md

    def repl(match: re.Match[str]) -> str:
        alt, target = match.group(1), match.group(2).strip()
        if re.match(r"^[a-z]+:", target, re.I) or target.startswith("#"):
            return match.group(0)
        src = (base / target).resolve()
        if not src.exists() or src.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}:
            return match.group(0)
        try:
            img = Image.open(src)
            if img.width <= max_width:
                return match.group(0)
            ratio = max_width / img.width
            img = img.resize((max_width, max(1, int(img.height * ratio))), Image.Resampling.LANCZOS)
            out = tmp / f"image-{len(list(tmp.glob('image-*'))) + 1:02d}.jpg"
            img.convert("RGB").save(out, "JPEG", quality=quality, optimize=True)
            return f"![{alt}]({out})"
        except Exception:
            return match.group(0)

    return re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", repl, md)


def set_doc_defaults(doc: Document, p: dict[str, str]) -> None:
    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.top_margin = Cm(2.15)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.25)
        section.right_margin = Cm(2.25)
        section.header_distance = Cm(1.1)
        section.footer_distance = Cm(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = p["font"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), p["font"])
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(p["text"])
    normal.paragraph_format.line_spacing = 1.33
    normal.paragraph_format.space_after = Pt(5)


def style_base(doc: Document, p: dict[str, str]) -> None:
    for name, size, before, after in [
        ("Title", 22, 0, 18),
        ("Heading 1", 18, 16, 10),
        ("Heading 2", 14, 18, 7),
        ("Heading 3", 11.5, 12, 5),
    ]:
        if name not in doc.styles:
            continue
        s = doc.styles[name]
        s.font.name = p["font"]
        s._element.rPr.rFonts.set(qn("w:eastAsia"), p["font"])
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = rgb(p["primary"] if name != "Heading 3" else p["accent"])
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
    for name in ("Source Code", "Code", "Verbatim Char"):
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = p["mono"]
            s._element.rPr.rFonts.set(qn("w:eastAsia"), p["mono"])
            s.font.size = Pt(8.8)


def make_reference_docx(path: Path, profile: str) -> None:
    default = subprocess.run([require("pandoc"), "--print-default-data-file", "reference.docx"], capture_output=True, check=False)
    if default.returncode != 0:
        die(default.stderr.decode("utf-8", "replace").strip() or "failed to create reference.docx")
    path.write_bytes(default.stdout)
    doc = Document(path)
    p = PALETTES[profile]
    set_doc_defaults(doc, p)
    style_base(doc, p)
    doc.save(path)


def set_cell_shading(cell, color_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def set_cell_text(cell, p: dict[str, str], bold=False, white=False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.15
        for run in para.runs:
            run.font.name = p["font"]
            run._element.rPr.rFonts.set(qn("w:eastAsia"), p["font"])
            run.font.size = Pt(9)
            run.font.bold = bold or run.font.bold
            run.font.color.rgb = RGBColor(255, 255, 255) if white else rgb(p["text"])


def cm_to_twips(cm: float) -> str:
    return str(int(cm * 567))


def set_table_grid(table, col_count: int, total_cm: float = 16.0) -> None:
    if col_count <= 0:
        return
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tbl_pr)
    for tag in ("w:tblW", "w:tblLayout"):
        old = tbl_pr.find(qn(tag))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), cm_to_twips(total_cm))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    col_w = cm_to_twips(total_cm / col_count)
    for _ in range(col_count):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), col_w)
        grid.append(col)
    tbl.insert(1, grid)


def set_cell_width(cell, cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcW"))
    if old is not None:
        tc_pr.remove(old)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), cm_to_twips(cm))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_table_border(table, color: str) -> None:
    tbl_pr = table._tbl.tblPr if table._tbl.tblPr is not None else OxmlElement("w:tblPr")
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)
    if table._tbl.tblPr is None:
        table._tbl.insert(0, tbl_pr)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def fit_table(table, p: dict[str, str]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_count = len(table.rows[0].cells) if table.rows else 0
    col_cm = 16.0 / max(col_count, 1)
    # ponytail: explicit tblGrid prevents LibreOffice/Google Docs from exploding markdown tables.
    set_table_grid(table, col_count)
    set_table_border(table, p["table_border"])
    for r, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        set_cant_split(row)
        for cell in row.cells:
            set_cell_width(cell, col_cm)
            if r == 0:
                set_cell_shading(cell, p["table_header_bg"])
                set_cell_text(cell, p, bold=True, white=True)
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if r % 2 == 0:
                    set_cell_shading(cell, p["table_alt_bg"])
                set_cell_text(cell, p)


def center_images(doc: Document) -> None:
    for para in doc.paragraphs:
        if para._p.xpath(".//pic:pic"):
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(8)
            para.paragraph_format.space_after = Pt(8)


def postprocess(docx_path: Path, profile: str, title: str | None) -> None:
    p = PALETTES[profile]
    doc = Document(docx_path)
    set_doc_defaults(doc, p)
    style_base(doc, p)

    if title and doc.paragraphs and doc.paragraphs[0].text.strip() != title:
        para = doc.paragraphs[0].insert_paragraph_before(title)
        para.style = doc.styles["Title"]

    for para in doc.paragraphs:
        name = para.style.name
        if name.startswith("Heading"):
            para.paragraph_format.keep_with_next = True
        if name in {"Source Code", "Code Block"}:
            para.paragraph_format.left_indent = Cm(0.25)
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.space_after = Pt(6)
        for run in para.runs:
            if run.text:
                font = p["mono"] if name in {"Source Code", "Code Block"} else p["font"]
                run.font.name = font
                run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
                if name in {"Title", "Heading 1", "Heading 2", "Heading 3"}:
                    run.font.color.rgb = rgb(p["primary"])
                    run.font.bold = True
    for table in doc.tables:
        fit_table(table, p)
    center_images(doc)
    doc.save(docx_path)

    # ponytail: opening media catches corrupt images cheaply; sizing is left to Word/Pandoc unless it breaks.
    with zipfile.ZipFile(docx_path) as z:
        for name in z.namelist():
            if name.startswith("word/media/"):
                with z.open(name) as f:
                    Image.open(f).verify()


def validate_docx(path: Path) -> None:
    doc = Document(path)
    if not doc.paragraphs and not doc.tables:
        die("empty document")


def convert(args: argparse.Namespace) -> None:
    md_path = Path(args.input).resolve()
    if not md_path.exists():
        die(f"input not found: {md_path}")
    out = Path(args.output).resolve() if args.output else md_path.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="md-docx-") as d:
        tmp = Path(d)
        tmp_md = tmp / "input.md"
        md = normalize_md(md_path.read_text(encoding="utf-8"))
        md = optimize_images(md, md_path.parent, tmp, args.max_image_width, args.image_quality)
        md = render_ascii_diagrams(render_mermaid(md, tmp), tmp)
        tmp_md.write_text(md, encoding="utf-8")
        ref = Path(args.reference_doc).resolve() if args.reference_doc else tmp / "reference.docx"
        if not ref.exists():
            make_reference_docx(ref, args.profile)
        cmd = [
            require("pandoc"), str(tmp_md),
            "--from", "markdown+smart+pipe_tables+grid_tables+multiline_tables+fenced_divs+bracketed_spans+link_attributes+yaml_metadata_block+tex_math_dollars+task_lists",
            "--to", "docx", "--standalone",
            "--reference-doc", str(ref),
            "--resource-path", f"{md_path.parent}:{tmp}:{Path.cwd()}",
            "--output", str(out),
        ]
        if args.toc:
            cmd += ["--toc", "--toc-depth", str(args.toc_depth)]
        if args.number_sections:
            cmd.append("--number-sections")
        p = run(cmd)
        if p.returncode != 0:
            die(p.stderr.strip() or "pandoc failed")

    postprocess(out, args.profile, args.title)
    validate_docx(out)
    print(out)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Convert Markdown to polished DOCX.")
    parser.add_argument("input", help="input .md")
    parser.add_argument("output", nargs="?", help="output .docx")
    parser.add_argument("--profile", choices=sorted(PALETTES), default="publish")
    parser.add_argument("--reference-doc", help="optional reference.docx")
    parser.add_argument("--title", help="title to insert if Markdown has no title")
    parser.add_argument("--toc", action="store_true", help="include Word TOC field")
    parser.add_argument("--toc-depth", type=int, default=3)
    parser.add_argument("--number-sections", action="store_true")
    parser.add_argument("--max-image-width", type=int, default=1600, help="resize raster images wider than this px; 0 disables")
    parser.add_argument("--image-quality", type=int, default=88, help="JPEG quality for resized images")
    return parser


if __name__ == "__main__":
    convert(build_parser().parse_args())
